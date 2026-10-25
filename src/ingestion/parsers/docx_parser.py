from pathlib import Path
from typing import Tuple
import docx

from src.ingestion.parsers.base import BaseParser
from src.schemas.ingest import DocumentMetadata


class DocxParser(BaseParser):
    """
    Parser for Word (.docx) Runemark documents using python-docx.
    """

    def parse(self, file_path: Path) -> Tuple[DocumentMetadata, str]:
        path = Path(file_path)
        if not path.is_file():
            raise FileNotFoundError(f"File not found: {file_path}")

        try:
            document = docx.Document(str(path))
        except Exception as e:
            raise ValueError(f"Corrupted or unreadable DOCX {path.name}: {e}")

        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        
        # In case there are tables with text
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n".join(paragraphs)
        if not full_text.strip():
            raise ValueError(f"DOCX {path.name} contains no text.")

        return self.extract_header_and_body(full_text, path, "DOCX")
